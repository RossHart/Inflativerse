import docx
import astropy.units as u
import datetime
import googlemaps
import math

# NOTE -- this code used the responses to the `old' form, and also only 
# includes the old managers. You will need to update this to use it.

name_lookups = {}
name_lookups['rh'] = {'name':'Ross Hart',
                      'office_number':'0115 846 8814',
                      'mobile_number':'07415 215 851'}

name_lookups['ra'] = {'name':'Rachel Asquith',
                      'office_number':'0115 846 62 81',
                      'mobile_number':'07891 528 045'}

name_lookups['ja'] = {'name':'Jake Arthur',
                      'office_number':'0115 846 64 00',
                      'mobile_number':'07964 068990'}

def load_document(document):
    
    doc = docx.Document(document)
    t = doc.tables[0]
    school_name = t.cell(1,1).text
    school_address = t.cell(2,1).text
    school_postcode = school_address[school_address.rfind('NG'):].replace(' ','')
    school_address = school_address[:school_address.rfind('NG')]
    school_postcode = school_postcode[:-3] + ' ' + school_postcode[-3:]
    contact_name = t.cell(3,1).text
    #start_time = t.cell(19,0).text
    return school_name, school_address, school_postcode, contact_name
  



def get_taxi_time(arrival_time,school_postcode,output='string',printout=False,
                  api_key = 'AIzaSyAI4IQJn00tsS34ZaHbBptWzeNG5OOZMyc'):

    university_postcode = 'NG7 2RD'
    # convert time to a datetime module format to work with:
    hm_separator = arrival_time.find(':')
    arrival_h = int(arrival_time[:hm_separator])
    arrival_m = int(arrival_time[hm_separator+1:])
    arrival_time = datetime.datetime(2000,1,1,arrival_h,arrival_m)

    client = googlemaps.client.Client(key=api_key)
    directions = googlemaps.directions.directions(client,university_postcode,
                                                  school_postcode,
                                                  arrival_time=arrival_time,
                                                  mode='driving')
    
    D = directions[0]['legs'][0]['distance']['value']*u.m
    t = directions[0]['legs'][0]['duration']['value']*u.s
    
    
    recommended_travel_time = t.to(u.minute) + 10*u.minute # + 10 mins for traffic/late taxis
    recommended_travel_time = math.ceil(recommended_travel_time.value/5)*5
    
    leaving_time = arrival_time - datetime.timedelta(0,60*recommended_travel_time)
    time_string = str(leaving_time.time())[:5]
    if printout is True:
        print('Google maps travel info: \n'+
              '------------------------ \n' + 
              'D={0:.2f}, '.format(D.to(u.km)) + 
              't={0:.2f}'.format(t.to(u.minute)))
        print('--> recommended = {} min'.format(recommended_travel_time))
        print('--> leave at {} \n------------------------'.format(time_string))
    if output is 'string':
        return time_string
    else:
        return leaving_time.time()


def taxi_email(document,arrival_time,end_time,date,your_initials,printout=False):
    
    school_name, school_address, school_postcode, _ = load_document(document)
    start_time = get_taxi_time(arrival_time,school_postcode,printout=printout)
    school_address = school_address.replace('\n',', ') + school_postcode
    university_address = 'Astronomy, University of Nottingham, NG7 2RD'
    
    preamble = ('Hi,\n\n'  
                'We would like to book a return taxi journey for {}. ' 
                'We would like to book the large (8 seater) airport taxis as we will also be ' 
                'transporting four large boxes. We have used the large airport taxis/vans before ' 
                'and the equipment and people fitted fine. The journeys (detailed below) will be ' 
                'between the astronomy building on University Park campus (NG7 2RD, building 25 on ' 
                'the attached map) and {} ({}).').format(date,school_name,school_postcode)
    
    def journey_box(pickup,destination,time,date,number=1):
        
        journey = ('\n\nJourney {}: {} - {}'
                   '\nPick up address: {}'
                   '\nDestination address: {}'
                   '\n5 people plus four large boxes').format(number,date,time,pickup,destination)
        return journey
    
    journey_1 = journey_box(university_address,school_address,start_time,date)
    journey_2 = journey_box(school_address,university_address,end_time,date,2)
    
    name = name_lookups[your_initials]['name']
    mobile_number = name_lookups[your_initials]['mobile_number']
    office_number = name_lookups[your_initials]['office_number']
    
    postamble = ('\n\nPlease charge the journeys to the following account: \n'
                 'Physics and Astronomy, project code A12707, under the name {}. The '
                 'contact telephone number is {} or {}. \n'
                 '\nMany thanks, \n{}').format(name,mobile_number,office_number,name)
    print(preamble,journey_1,journey_2,postamble)
    return None

document = input('Enter document name: ')
arrival_time = input('Enter arrival time: ')
end_time = input('Enter end time: ')
date = input('Enter date of trip: ')
your_initials = input('Enter your initials (ja/ra/rh): ')
printout = input('Print googlemaps journey info? (y/n): ')
if printout is 'y':
    printout = True
else:
    printout = False

taxi_email(document,arrival_time,end_time,date,your_initials,printout)


